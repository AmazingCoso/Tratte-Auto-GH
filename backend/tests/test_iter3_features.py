# Iteration 3: suggestions, per-leg dates in Excel, trips storing legs/km, monthly report
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
MI, BO, FI = "45.4642,9.19", "44.4949,11.3426", "43.7696,11.2558"
MONTH = "2031-03"


@pytest.fixture(scope="module")
def s(profile_id):
    sess = requests.Session()
    sess.headers["X-Profile-Id"] = profile_id
    yield sess
    for t in sess.get(f"{BASE_URL}/api/trips", params={"q": "TEST_"}).json():
        if t["name"].startswith("TEST_"):
            sess.delete(f"{BASE_URL}/api/trips/{t['id']}")


def wb_from(r):
    return load_workbook(io.BytesIO(r.content)).active


# --- Address suggestions
def test_suggest_short_query_empty(s):
    r = s.get(f"{BASE_URL}/api/route/suggest", params={"q": "Mi"})
    assert r.status_code == 200 and r.json() == []


def test_suggest_returns_strings(s):
    r = s.get(f"{BASE_URL}/api/route/suggest", params={"q": "Via Roma Milano"})
    assert r.status_code == 200
    data = r.json()
    assert isinstance(data, list) and len(data) > 0
    assert all(isinstance(x, str) and x for x in data)
    assert len(data) == len(set(data))


# --- Per-leg dates in leg Excel
def test_leg_excel_per_leg_dates(s):
    r = s.post(f"{BASE_URL}/api/route/excel", json={
        "origin": MI, "destination": FI, "waypoints": [BO], "mode": "driving",
        "date": "2031-03-01", "leg_dates": ["2031-03-01", "2031-03-05"], "name": "TEST_x"})
    assert r.status_code == 200
    ws = wb_from(r)
    assert [c.value for c in ws[4]][1] == "Data"
    d5, d6 = ws.cell(5, 2).value, ws.cell(6, 2).value
    assert d5.strftime("%Y-%m-%d") == "2031-03-01"
    assert d6.strftime("%Y-%m-%d") == "2031-03-05"


def test_leg_excel_fallback_trip_date(s):
    r = s.post(f"{BASE_URL}/api/route/excel", json={
        "origin": MI, "destination": FI, "waypoints": [BO], "date": "2031-03-02", "leg_dates": ["2031-03-04"]})
    ws = wb_from(r)
    assert ws.cell(5, 2).value.strftime("%Y-%m-%d") == "2031-03-04"
    assert ws.cell(6, 2).value.strftime("%Y-%m-%d") == "2031-03-02"


# --- Trips store date/leg_dates/legs/total_km
@pytest.fixture(scope="module")
def trip(s):
    r = s.post(f"{BASE_URL}/api/trips", json={
        "name": "TEST_iter3", "origin": MI, "destination": FI, "waypoints": [BO], "mode": "driving",
        "date": "2031-03-10", "leg_dates": ["2031-03-10", "2031-04-02"]})
    assert r.status_code == 200, r.text
    return r.json()


def test_trip_saves_legs_and_km(s, trip):
    assert trip["date"] == "2031-03-10"
    assert trip["leg_dates"] == ["2031-03-10", "2031-04-02"]
    assert isinstance(trip["legs"], list) and len(trip["legs"]) == 2
    assert trip["total_km"] > 250
    assert trip["total_min"] > 100
    g = s.get(f"{BASE_URL}/api/trips/{trip['id']}").json()
    assert g["legs"] == trip["legs"] and g["total_km"] == trip["total_km"]
    assert "_id" not in g


def test_trip_invalid_date_422(s):
    r = s.post(f"{BASE_URL}/api/trips", json={"name": "TEST_bad", "origin": MI, "destination": FI, "date": "10/03/2031"})
    assert r.status_code == 422


def test_trip_default_date(s):
    r = s.post(f"{BASE_URL}/api/trips", json={"name": "TEST_nodate", "origin": MI, "destination": BO})
    assert r.status_code == 200
    assert len(r.json()["date"]) == 10 and r.json()["total_km"] > 0


def test_trip_unresolvable_saves_null_legs(s):
    r = s.post(f"{BASE_URL}/api/trips", json={"name": "TEST_unres", "origin": "zzqxw-nonexistent-xyz123", "destination": MI})
    assert r.status_code == 200
    assert r.json()["legs"] is None and r.json()["total_km"] is None


def test_update_recomputes_legs(s, trip):
    r = s.put(f"{BASE_URL}/api/trips/{trip['id']}", json={
        "name": "TEST_iter3", "origin": MI, "destination": FI, "waypoints": [BO], "mode": "driving",
        "date": "2031-03-10", "leg_dates": ["2031-03-10", "2031-04-02"]})
    assert r.status_code == 200 and len(r.json()["legs"]) == 2


def test_duplicate_copies_legs(s, trip):
    r = s.post(f"{BASE_URL}/api/trips/{trip['id']}/duplicate")
    d = r.json()
    assert d["name"] == "TEST_iter3 (copia)"
    assert d["legs"] == trip["legs"] and d["leg_dates"] == trip["leg_dates"] and d["date"] == trip["date"]
    s.delete(f"{BASE_URL}/api/trips/{d['id']}")


# --- Monthly report
def test_monthly_report_per_leg_dates(s, trip):
    r = s.get(f"{BASE_URL}/api/reports/monthly", params={"month": MONTH})
    assert r.status_code == 200
    d = r.json()
    assert d["legs"] == 1 and d["trips"] == 1
    assert d["total_km"] == round(trip["legs"][0]["distance_km"], 2)
    a = s.get(f"{BASE_URL}/api/reports/monthly", params={"month": "2031-04"}).json()
    assert a["legs"] == 1 and a["total_km"] == round(trip["legs"][1]["distance_km"], 2)


def test_monthly_report_bad_month_422(s):
    assert s.get(f"{BASE_URL}/api/reports/monthly", params={"month": "bad"}).status_code == 422
    assert s.get(f"{BASE_URL}/api/reports/monthly/excel", params={"month": "2031-3"}).status_code == 422


def test_monthly_report_empty_month(s):
    d = s.get(f"{BASE_URL}/api/reports/monthly", params={"month": "1999-01"}).json()
    assert d == {"month": "1999-01", "trips": 0, "legs": 0, "total_km": 0, "total_min": 0}
    r = s.get(f"{BASE_URL}/api/reports/monthly/excel", params={"month": "1999-01"})
    assert r.status_code == 200


def test_monthly_excel_content(s, trip):
    r = s.get(f"{BASE_URL}/api/reports/monthly/excel", params={"month": MONTH})
    assert r.status_code == 200
    assert 'report_km_2031-03.xlsx' in r.headers.get("content-disposition", "")
    ws = wb_from(r)
    assert [c.value for c in ws[4]] == ["Data", "Viaggio", "Percorso", "Distanza (km)"]
    assert ws.cell(5, 2).value == "TEST_iter3"
    assert ws.cell(5, 1).value.strftime("%Y-%m-%d") == "2031-03-10"
    last = ws.max_row
    assert ws.cell(last, 3).value == "TOTALE"
    assert str(ws.cell(last, 4).value).startswith("=SUM(D5:D")
