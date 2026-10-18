import asyncio
import io
import re
from datetime import date as Date
from typing import List, Literal, Optional

import httpx
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel, Field

router = APIRouter(prefix="/api/route")

UA = {"User-Agent": "RutaPianificatoreViaggio/1.0"}
PROFILES = {"driving": "routed-car", "walking": "routed-foot", "bicycling": "routed-bike", "transit": "routed-car"}
COORD_RX = re.compile(r"^\s*(-?\d{1,2}(?:\.\d+)?)\s*,\s*(-?\d{1,3}(?:\.\d+)?)\s*$")
MODE_LABELS = {"driving": "Auto", "walking": "A piedi", "bicycling": "In bici", "transit": "Mezzi pubblici (stima su strada)"}
_geo_cache: dict = {}
_geo_lock = asyncio.Lock()


class RouteInput(BaseModel):
    origin: str = Field(min_length=1)
    destination: str = Field(min_length=1)
    waypoints: List[str] = []
    mode: Literal["driving", "walking", "bicycling", "transit"] = "driving"


class ExcelInput(RouteInput):
    date: Date
    leg_dates: List[Date] = []
    name: Optional[str] = None


async def geocode(client: httpx.AsyncClient, query: str):
    m = COORD_RX.match(query)
    if m:
        return float(m.group(1)), float(m.group(2))
    key = query.strip().lower()
    if key in _geo_cache:
        return _geo_cache[key]
    async with _geo_lock:  # Nominatim policy: max 1 req/s
        r = await client.get(
            "https://nominatim.openstreetmap.org/search",
            params={"q": query, "format": "json", "limit": 1, "accept-language": "it"},
        )
        await asyncio.sleep(1)
    data = r.json() if r.status_code == 200 else []
    if not data:
        raise HTTPException(status_code=422, detail=f"Indirizzo non trovato: {query}")
    _geo_cache[key] = (float(data[0]["lat"]), float(data[0]["lon"]))
    return _geo_cache[key]


async def compute_legs(data: RouteInput) -> dict:
    stops = [data.origin.strip(), *[w.strip() for w in data.waypoints if w.strip()], data.destination.strip()]
    async with httpx.AsyncClient(headers=UA, timeout=20) as client:
        coords = [await geocode(client, s) for s in stops]
        path = ";".join(f"{lon},{lat}" for lat, lon in coords)
        r = await client.get(
            f"https://routing.openstreetmap.de/{PROFILES[data.mode]}/route/v1/driving/{path}",
            params={"overview": "false"},
        )
    body = r.json() if r.status_code == 200 else {}
    if body.get("code") != "Ok":
        raise HTTPException(status_code=422, detail="Impossibile calcolare il percorso tra queste località")
    legs = [
        {
            "index": i + 1,
            "from": stops[i],
            "to": stops[i + 1],
            "distance_km": round(leg["distance"] / 1000, 2),
            "duration_min": round(leg["duration"] / 60, 1),
        }
        for i, leg in enumerate(body["routes"][0]["legs"])
    ]
    return {
        "legs": legs,
        "total_km": round(sum(l["distance_km"] for l in legs), 2),
        "total_min": round(sum(l["duration_min"] for l in legs), 1),
        "mode": data.mode,
    }


def fmt_duration(minutes: float) -> str:
    h, m = divmod(int(round(minutes)), 60)
    return f"{h} h {m:02d} min" if h else f"{m} min"


@router.post("/legs")
async def route_legs(data: RouteInput):
    return await compute_legs(data)


LEG_HEADERS = ["N°", "Data", "Partenza", "Arrivo", "Distanza (km)", "Durata (min)", "Durata"]


def write_xlsx(title: str, subtitle: str, headers: list, rows: list, total_min: float, filename: str):
    """rows must end with (..., km, min, durata) columns; TOTALE row sums the km/min columns."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tratte"
    thin = Side(style="thin", color="E2E8F0")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws.append([])
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F172A")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))

    n = len(headers)
    km_col, min_col = ws.cell(4, n - 2).column_letter, ws.cell(4, n - 1).column_letter
    last = max(ws.max_row, 5)
    total = [""] * n
    total[n - 4] = "TOTALE"
    total[n - 3] = f"=SUM({km_col}5:{km_col}{last})"
    total[n - 2] = f"=SUM({min_col}5:{min_col}{last})"
    total[n - 1] = fmt_duration(total_min)
    ws.append(total)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F1F5F9")

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if isinstance(cell.value, Date):
                cell.number_format = "DD/MM/YYYY"
        row[n - 3].number_format = "0.00"
        row[n - 2].number_format = "0.0"

    for i in range(1, n + 1):
        h = headers[i - 1]
        ws.column_dimensions[ws.cell(4, i).column_letter].width = 36 if h in ("Partenza", "Arrivo", "Viaggio") else max(8, len(h) + 4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/suggest")
async def suggest(q: str):
    if len(q.strip()) < 3:
        return []
    async with httpx.AsyncClient(headers=UA, timeout=8) as client:
        r = await client.get("https://photon.komoot.io/api/", params={"q": q, "limit": 5, "lat": 42.5, "lon": 12.5})
    if r.status_code != 200:
        return []
    out = []
    for f in r.json().get("features", []):
        p = f["properties"]
        street = " ".join(x for x in [p.get("street"), p.get("housenumber")] if x)
        city = " ".join(x for x in [p.get("postcode"), p.get("city") or p.get("county")] if x)
        parts = [p.get("name"), street, city, p.get("country")]
        label = ", ".join(dict.fromkeys(x for x in parts if x))
        if label and label not in out:
            out.append(label)
    return out


@router.post("/excel")
async def route_excel(data: ExcelInput):
    result = await compute_legs(data)
    rows = [
        [l["index"], data.leg_dates[i] if i < len(data.leg_dates) else data.date, l["from"], l["to"],
         l["distance_km"], l["duration_min"], fmt_duration(l["duration_min"])]
        for i, l in enumerate(result["legs"])
    ]
    return write_xlsx(
        data.name or "Viaggio", f"Mezzo: {MODE_LABELS[data.mode]}", LEG_HEADERS, rows,
        result["total_min"], f"tratte_{data.date.isoformat()}.xlsx",
    )
