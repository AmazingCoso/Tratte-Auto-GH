"""Backend tests for /api/route/legs and /api/route/excel (Nominatim + OSRM)."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://trattative-team-3.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def s():
    sess = requests.Session()
    sess.headers.update({"Content-Type": "application/json"})
    return sess


# ----- /api/route/legs -----
def test_legs_coords_only(s):
    """Pure coord input bypasses Nominatim (fast path)."""
    # Milano ~ Bologna coords
    r = s.post(f"{API}/route/legs", json={
        "origin": "45.4642,9.1900",
        "destination": "44.4949,11.3426",
        "waypoints": [],
        "mode": "driving",
    }, timeout=60)
    assert r.status_code == 200, r.text
    data = r.json()
    assert "legs" in data and len(data["legs"]) == 1
    leg = data["legs"][0]
    assert leg["index"] == 1
    assert leg["distance_km"] > 150 and leg["distance_km"] < 300
    assert leg["duration_min"] > 60
    assert data["total_km"] == leg["distance_km"]
    assert data["total_min"] == leg["duration_min"]
    assert data["mode"] == "driving"


def test_legs_with_waypoint_cities(s):
    """Nominatim path with cached Italian cities, multi-leg."""
    r = s.post(f"{API}/route/legs", json={
        "origin": "Milano",
        "destination": "Firenze",
        "waypoints": ["Bologna"],
        "mode": "driving",
    }, timeout=120)
    assert r.status_code == 200, r.text
    data = r.json()
    assert len(data["legs"]) == 2
    assert data["legs"][0]["from"].lower().startswith("milano")
    assert data["legs"][0]["to"].lower().startswith("bologna")
    assert data["legs"][1]["from"].lower().startswith("bologna")
    assert data["legs"][1]["to"].lower().startswith("firenze")
    # totals = sum of legs (allow tiny rounding)
    s_km = round(sum(l["distance_km"] for l in data["legs"]), 2)
    assert abs(data["total_km"] - s_km) < 0.1


def test_legs_unknown_address_422(s):
    r = s.post(f"{API}/route/legs", json={
        "origin": "Milano",
        "destination": "zzzzz-nonexistent-place-xyzxyz-qqq",
        "waypoints": [],
        "mode": "driving",
    }, timeout=60)
    assert r.status_code == 422
    assert "Indirizzo non trovato" in r.json()["detail"]


def test_legs_missing_fields_422(s):
    r = s.post(f"{API}/route/legs", json={"origin": "", "destination": "Roma", "mode": "driving"})
    assert r.status_code == 422


# ----- /api/route/excel -----
def test_excel_download_and_content(s):
    r = s.post(f"{API}/route/excel", json={
        "origin": "45.4642,9.1900",
        "destination": "44.4949,11.3426",
        "waypoints": [],
        "mode": "driving",
        "date": "2026-01-15",
        "name": "TEST_Excel Trip",
    }, timeout=60)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "tratte_2026-01-15.xlsx" in r.headers.get("content-disposition", "")

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    # Row 1 title, row 2 mode, row 4 headers, row 5 first leg, last row TOTALE
    assert ws["A1"].value == "TEST_Excel Trip"
    assert "Mezzo" in str(ws["A2"].value)
    headers = [c.value for c in ws[4]]
    assert headers == ["N°", "Data", "Partenza", "Arrivo", "Distanza (km)", "Durata (min)", "Durata"]
    # TOTALE row at bottom
    last_row = [c.value for c in ws[ws.max_row]]
    assert "TOTALE" in last_row
    # SUM formulas
    assert str(last_row[4]).startswith("=SUM(E5:E")
    assert str(last_row[5]).startswith("=SUM(F5:F")


def test_excel_invalid_date_422(s):
    r = s.post(f"{API}/route/excel", json={
        "origin": "Milano",
        "destination": "Roma",
        "waypoints": [],
        "mode": "driving",
        "date": "not-a-date",
    })
    assert r.status_code == 422
