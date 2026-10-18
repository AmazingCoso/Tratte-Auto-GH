# Iteration 10: regression for write_xlsx refactor (headers, TOTALE row, formulas, formats, widths)
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
MONTH = "2031-04"


@pytest.fixture(scope="module")
def s(profile_id):
    sess = requests.Session()
    sess.headers.update({"X-Profile-Id": profile_id})
    return sess


def _wb(r):
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    return load_workbook(io.BytesIO(r.content)).active


def test_route_excel_structure(s):
    r = s.post(f"{BASE_URL}/api/route/excel", json={
        "origin": "45.4642,9.19", "destination": "45.0703,7.6869", "waypoints": ["44.4056,8.9463"],
        "mode": "driving", "date": "2031-04-10", "leg_dates": ["2031-04-10", "2031-04-11"], "name": "TEST_xls"})
    assert 'filename="tratte_2031-04-10.xlsx"' in r.headers["content-disposition"]
    ws = _wb(r)
    assert ws["A1"].value == "TEST_xls" and ws["A1"].font.bold
    assert ws["A2"].value == "Mezzo: Auto"
    assert [c.value for c in ws[4]] == ["N°", "Data", "Partenza", "Arrivo", "Distanza (km)", "Durata (min)", "Durata"]
    assert ws["A4"].fill.fgColor.rgb.endswith("0F172A")
    assert ws.max_row == 7  # header + 2 legs + total
    assert ws["A5"].value == 1 and ws["A6"].value == 2
    assert ws["B5"].number_format == "DD/MM/YYYY" and ws["B6"].value.day == 11
    assert ws["E5"].number_format == "0.00" and ws["F5"].number_format == "0.0"
    assert ws["D7"].value == "TOTALE"
    assert ws["E7"].value == "=SUM(E5:E6)" and ws["F7"].value == "=SUM(F5:F6)"
    assert "min" in ws["G7"].value and ws["G7"].font.bold
    assert ws.column_dimensions["C"].width == 36 and ws.column_dimensions["A"].width == 8


def test_monthly_excel_structure(s):
    t = s.post(f"{BASE_URL}/api/trips", json={
        "name": "TEST_xls_month", "origin": "45.4642,9.19", "destination": "45.0703,7.6869",
        "waypoints": [], "mode": "driving", "date": f"{MONTH}-05", "leg_dates": []})
    assert t.status_code == 200
    try:
        ws = _wb(s.get(f"{BASE_URL}/api/reports/monthly/excel", params={"month": MONTH}))
        assert [c.value for c in ws[4]] == ["Data", "Viaggio", "Percorso", "Distanza (km)"]
        assert ws["B5"].value == "TEST_xls_month"
        last = ws.max_row
        assert ws[f"C{last}"].value == "TOTALE"
        assert ws[f"D{last}"].value == f"=SUM(D5:D{last - 1})"
        assert ws.column_dimensions["C"].width == 70 and ws.column_dimensions["B"].width == 30
    finally:
        s.delete(f"{BASE_URL}/api/trips/{t.json()['id']}")


def test_monthly_excel_empty_month(s):
    ws = _wb(s.get(f"{BASE_URL}/api/reports/monthly/excel", params={"month": "1999-01"}))
    assert ws["C5"].value == "TOTALE" and ws["D5"].value == "=SUM(D5:D5)"
